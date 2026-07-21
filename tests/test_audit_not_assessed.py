"""Acceptance: with optional sources OFF, the audit never records a false pass.

Every RENDER and EXTERNAL check, and every CWV check that has no static-heuristic
part, must report exactly "Not assessed" — never Present/Not present.
"""

import pandas as pd
import pytest

from advertools_mcp.audit.catalogue import load_catalogue
from advertools_mcp.audit.engine import run_audit
from advertools_mcp.config import Settings

# CWV checks that have a legitimate static-heuristic part (run regardless).
CWV_PARTIAL = {28, 29, 53, 85, 86}
ALLOWED_NON_ASSESSED_STATES = {"Not assessed"}
ALLOWED_HEURISTIC_STATES = {"Not assessed", "Not present", "Heuristic"}


def _toy_crawl(tmp_path):
    df = pd.DataFrame(
        {
            "url": ["https://site.test/a", "https://site.test/b"],
            "status": [200, 200],
            "title": ["A", "B"],
            "meta_desc": ["da", "db"],
            "meta_robots": ["index,follow", "noindex"],
            "h1": ["A", "B"],
            "canonical": ["https://site.test/a", "https://site.test/b"],
            "body_text": ["hello world " * 30, "more text " * 30],
            "resp_headers_Content-Type": ["text/html", "text/html"],
            "img_src": ["https://site.test/i.png", ""],
            "img_alt": ["alt", ""],
            "links_url": ["https://site.test/b", "https://site.test/a"],
            "links_nofollow": ["False", "False"],
        }
    )
    p = tmp_path / "crawl.parquet"
    df.to_parquet(p, index=False)
    return str(p)


@pytest.fixture
def audit_summary(tmp_path):
    parquet = _toy_crawl(tmp_path)
    # All optional sources OFF.
    settings = Settings(
        data_dir=tmp_path,
        lighthouse_api_key="",
        gsc_credentials="",
        enable_render=False,
    )
    settings.ensure_dirs()
    run_audit(parquet, settings, settings.audits_dir, has_backlinks=False)
    # Read back the per-check statuses from the workbook.
    from openpyxl import load_workbook

    audits = list(settings.audits_dir.glob("*.xlsx"))
    wb = load_workbook(audits[0], read_only=True, data_only=True)
    ws = wb["Checklist"]
    statuses = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or row[0] is None:
            continue
        statuses[int(row[0])] = row[3]
    wb.close()
    return statuses


def test_render_checks_not_assessed(audit_summary):
    for c in load_catalogue():
        if c.tier == "RENDER":
            assert audit_summary[c.num] == "Not assessed", f"#{c.num} {c.check}"


def test_external_checks_not_assessed(audit_summary):
    for c in load_catalogue():
        if c.tier == "EXTERNAL":
            assert audit_summary[c.num] == "Not assessed", f"#{c.num} {c.check}"


def test_cwv_pure_checks_not_assessed(audit_summary):
    for c in load_catalogue():
        if c.tier == "CWV" and c.num not in CWV_PARTIAL:
            assert audit_summary[c.num] == "Not assessed", f"#{c.num} {c.check}"


def test_cwv_partial_checks_never_false_pass(audit_summary):
    # Heuristic CWV checks may be Heuristic/Not present, but never a hard Present
    # without the heuristic label, and never a fabricated pass.
    for c in load_catalogue():
        if c.tier == "CWV" and c.num in CWV_PARTIAL:
            assert audit_summary[c.num] in ALLOWED_HEURISTIC_STATES, f"#{c.num} {c.check}"


def test_no_check_uses_forbidden_state(audit_summary):
    allowed = {"Present", "Not present", "Not assessed", "Heuristic"}
    for num, status in audit_summary.items():
        assert status in allowed, f"#{num} has illegal state {status!r}"


def test_workbook_has_descriptions_and_crawl_detail_sheet(tmp_path):
    from openpyxl import load_workbook

    parquet = _toy_crawl(tmp_path)
    settings = Settings(data_dir=tmp_path)
    settings.ensure_dirs()
    run_audit(parquet, settings, settings.audits_dir)
    wb = load_workbook(list(settings.audits_dir.glob("*.xlsx"))[0], read_only=True, data_only=True)

    # Crawl Detail is embedded as its own sheet.
    assert "Crawl Detail" in wb.sheetnames
    cd = wb["Crawl Detail"]
    header = [c.value for c in next(cd.iter_rows(max_row=1))]
    assert header[0] == "Address"

    # Every check row carries a plain-English explanation (column E).
    ws = wb["Checklist"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header[4] == "What it checks & why it matters"
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert all((r[4] or "").strip() for r in rows if r[0] is not None)
    wb.close()
