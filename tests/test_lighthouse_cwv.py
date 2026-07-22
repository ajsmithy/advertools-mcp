"""CWV tier: PageSpeed Insights integration (mocked — no real API calls)."""

import pandas as pd
import pytest

from advertools_mcp.audit import lighthouse
from advertools_mcp.audit.checks import CHECKS
from advertools_mcp.audit.context import AuditContext
from advertools_mcp.audit.engine import run_audit
from advertools_mcp.config import Settings


def _ctx(psi_results, has_lighthouse=True):
    df = pd.DataFrame({"url": ["https://s.test/"], "status": [200]})
    return AuditContext(
        df=df, settings=Settings(), has_lighthouse=has_lighthouse, psi_results=psi_results
    )


def _psi(url="https://s.test/", **audits):
    return {url: {"audits": audits}}


def test_lcp_present_when_slow():
    ctx = _ctx(_psi(**{"largest-contentful-paint": {"score": 0.3, "numericValue": 4800}}))
    res = CHECKS[25](ctx)
    assert res.status == "Present"
    assert res.affected_urls == ["https://s.test/"]
    assert res.detection_source == "PageSpeed Insights"


def test_lcp_not_present_when_fast():
    ctx = _ctx(_psi(**{"largest-contentful-paint": {"score": 0.95, "numericValue": 1400}}))
    assert CHECKS[25](ctx).status == "Not present"


def test_cls_thresholds():
    bad = _ctx(_psi(**{"cumulative-layout-shift": {"score": 0.4, "numericValue": 0.31}}))
    ok = _ctx(_psi(**{"cumulative-layout-shift": {"score": 1.0, "numericValue": 0.02}}))
    assert CHECKS[26](bad).status == "Present"
    assert CHECKS[26](ok).status == "Not present"


def test_binary_audit_score():
    bad = _ctx(_psi(**{"third-party-facades": {"score": 0, "numericValue": None}}))
    ok = _ctx(_psi(**{"third-party-facades": {"score": 1, "numericValue": None}}))
    assert CHECKS[88](bad).status == "Present"
    assert CHECKS[88](ok).status == "Not present"


def test_renamed_lcp_preload_audit_both_names():
    new = _ctx(_psi(**{"prioritize-lcp-image": {"score": 0, "numericValue": None}}))
    old = _ctx(_psi(**{"preload-lcp-image": {"score": 0, "numericValue": None}}))
    assert CHECKS[87](new).status == "Present"
    assert CHECKS[87](old).status == "Present"


def test_partial_check_upgrades_from_heuristic_to_real():
    ctx = _ctx(_psi(**{"render-blocking-resources": {"score": 0.2, "numericValue": 900}}))
    res = CHECKS[28](ctx)
    assert res.status == "Present"  # real verdict, not Heuristic
    assert res.detection_source == "PageSpeed Insights"


def test_missing_audit_falls_back_not_false_verdict():
    # PSI resolved but this Lighthouse version lacks uses-rel-preload:
    # check 86 must fall back to its heuristic, never fabricate a PSI verdict.
    ctx = _ctx(_psi(**{"largest-contentful-paint": {"score": 1.0, "numericValue": 900}}))
    res = CHECKS[86](ctx)
    assert res.detection_source != "PageSpeed Insights"


def test_no_key_still_not_assessed():
    ctx = _ctx(None, has_lighthouse=False)
    for num in (25, 26, 27, 87, 88):
        assert CHECKS[num](ctx).status == "Not assessed"


def test_key_but_psi_failed_is_not_assessed():
    ctx = _ctx(None, has_lighthouse=True)
    res = CHECKS[25](ctx)
    assert res.status == "Not assessed"
    assert "PageSpeed" in res.note or "PageSpeed" in res.detection_source


def test_engine_wires_psi_results(tmp_path, monkeypatch):
    df = pd.DataFrame({
        "url": ["https://s.test/"], "status": [200],
        "resp_headers_Content-Type": ["text/html"], "depth": [0],
    })
    pq = tmp_path / "c.parquet"
    df.to_parquet(pq, index=False)

    calls = {}

    def fake_psi(urls, key, strategy="mobile"):
        calls["urls"], calls["key"], calls["strategy"] = list(urls), key, strategy
        return {u: {"audits": {"largest-contentful-paint": {"score": 0.2, "numericValue": 5000}}}
                for u in urls}

    monkeypatch.setattr(lighthouse, "run_psi_sample", fake_psi)
    settings = Settings(data_dir=tmp_path, lighthouse_api_key="k123")
    settings.ensure_dirs()
    res = run_audit(str(pq), settings, settings.audits_dir)

    assert calls["key"] == "k123"
    assert calls["urls"] == ["https://s.test/"]
    # LCP (#25) must now be a real finding, and PSI config recorded.
    from openpyxl import load_workbook
    wb = load_workbook(res["audit_xlsx"], read_only=True, data_only=True)
    statuses = {int(r[0]): r[3] for i, r in enumerate(wb["Checklist"].iter_rows(values_only=True))
                if i > 0 and r[0] is not None}
    assert statuses[25] == "Present"
    assert res["config"].get("PSI URLs assessed") == "1"
    wb.close()


def test_run_psi_sample_parses_and_filters(monkeypatch):
    class FakeResp:
        def raise_for_status(self): pass
        def json(self):
            return {"lighthouseResult": {"audits": {
                "largest-contentful-paint": {"score": 0.5, "numericValue": 3000, "title": "x"},
                "unrelated-audit": {"score": 1},
            }}}

    class FakeClient:
        def __init__(self, **kw): pass
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def get(self, url, params=None):
            assert params["key"] == "k"
            return FakeResp()

    monkeypatch.setattr(lighthouse.httpx, "Client", FakeClient)
    out = lighthouse.run_psi_sample(["https://s.test/"], "k")
    audits = out["https://s.test/"]["audits"]
    assert "largest-contentful-paint" in audits
    assert "unrelated-audit" not in audits  # only catalogue-relevant audits kept
