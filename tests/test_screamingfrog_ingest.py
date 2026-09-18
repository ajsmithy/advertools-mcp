"""Auditing an imported Screaming Frog export folder."""

import pandas as pd
import pytest

from advertools_mcp.audit.engine import run_audit_from_folder
from advertools_mcp.audit.ingest import ingest_folder
from advertools_mcp.config import Settings

# Screaming Frog Internal:All style rows (real SF column names).
INTERNAL_ROWS = [
    {"Address": "https://s.test/", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Indexable",
     "Title 1": "Home", "H1-1": "Welcome", "Meta Description 1": "The home page.",
     "Canonical Link Element 1": "https://s.test/", "Meta Robots 1": "index,follow",
     "Word Count": "300", "Crawl Depth": "0", "Language": "en"},
    {"Address": "https://s.test/no-h1", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Indexable",
     "Title 1": "No H1 Page", "H1-1": "", "Meta Description 1": "d",
     "Canonical Link Element 1": "https://s.test/no-h1", "Meta Robots 1": "index,follow",
     "Word Count": "120", "Crawl Depth": "1", "Language": "en"},
    {"Address": "https://s.test/no-canonical", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Indexable",
     "Title 1": "No Canon", "H1-1": "Hi", "Meta Description 1": "d",
     "Canonical Link Element 1": "", "Meta Robots 1": "index,follow",
     "Word Count": "90", "Crawl Depth": "1", "Language": "en"},
    {"Address": "https://s.test/multi-title", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Indexable",
     "Title 1": "First", "Title 2": "Second", "H1-1": "X",
     "Meta Description 1": "d", "Canonical Link Element 1": "https://s.test/multi-title",
     "Meta Robots 1": "index,follow", "Word Count": "200", "Crawl Depth": "1", "Language": "en"},
    {"Address": "https://s.test/multi-desc", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Indexable",
     "Title 1": "Desc", "H1-1": "Y", "Meta Description 1": "a", "Meta Description 2": "b",
     "Canonical Link Element 1": "https://s.test/multi-desc", "Meta Robots 1": "index,follow",
     "Word Count": "200", "Crawl Depth": "1", "Language": "en"},
    {"Address": "https://s.test/noindex", "Content Type": "text/html; charset=UTF-8",
     "Status Code": "200", "Status": "OK", "Indexability": "Non-Indexable",
     "Title 1": "Noindex", "H1-1": "Z", "Meta Description 1": "d",
     "Canonical Link Element 1": "https://s.test/noindex",
     "Meta Robots 1": "noindex,nosnippet,noarchive",
     "Word Count": "200", "Crawl Depth": "1", "Language": "en"},
    {"Address": "https://s.test/old", "Content Type": "text/html",
     "Status Code": "301", "Status": "Moved Permanently", "Indexability": "Non-Indexable",
     "Title 1": "", "H1-1": "", "Canonical Link Element 1": "",
     "Redirect URL": "https://s.test/new", "Redirect Type": "HTTP Redirect",
     "Crawl Depth": "1"},
]

OUTLINKS_ROWS = [
    {"Type": "Hyperlink", "Source": "https://s.test/", "Destination": "https://s.test/no-h1",
     "Anchor": "No H1", "Follow": "true", "Status Code": "200"},
    {"Type": "Hyperlink", "Source": "https://s.test/",
     "Destination": "https://s.test/search?q=x&utm_source=nav", "Anchor": "Search",
     "Follow": "true", "Status Code": "200"},
    {"Type": "JS", "Source": "https://s.test/", "Destination": "https://cdn.ext/app.js",
     "Anchor": "", "Follow": "true", "Status Code": "200"},
    {"Type": "CSS", "Source": "https://s.test/", "Destination": "https://cdn.ext/style.css",
     "Anchor": "", "Follow": "true", "Status Code": "200"},
    {"Type": "Hyperlink", "Source": "https://s.test/no-h1", "Destination": "https://s.test/",
     "Anchor": "Home", "Follow": "false", "Status Code": "200"},
]


def _write_internal(folder):
    pd.DataFrame(INTERNAL_ROWS).to_csv(folder / "internal_all.csv", index=False)


def _write_outlinks(folder):
    pd.DataFrame(OUTLINKS_ROWS).to_csv(folder / "all_outlinks.csv", index=False)


def _status_map(summary):
    from openpyxl import load_workbook

    wb = load_workbook(summary["audit_xlsx"], read_only=True, data_only=True)
    out = {}
    for i, row in enumerate(wb["Checklist"].iter_rows(values_only=True)):
        if i == 0 or row[0] is None:
            continue
        out[int(row[0])] = row[3]
    wb.close()
    return out


@pytest.fixture
def settings(tmp_path):
    s = Settings(data_dir=tmp_path / "data")
    s.ensure_dirs()
    return s


def test_ingest_maps_columns_and_signals(tmp_path):
    _write_internal(tmp_path)
    _write_outlinks(tmp_path)
    res = ingest_folder(str(tmp_path))
    assert res.row_count == 7
    assert res.source == "screamingfrog"
    assert {"title", "canonical", "meta_robots", "multiples", "links", "scripts"} <= res.available_signals
    assert "viewport" not in res.available_signals
    assert "structured_data" not in res.available_signals
    # Multiples counts populated for #56/#76.
    row = res.df[res.df["url"] == "https://s.test/multi-title"].iloc[0]
    assert row["title_count"] == 2
    # Link graph reconstructed from All Outlinks.
    home = res.df[res.df["url"] == "https://s.test/"].iloc[0]
    assert "utm_source" in home["links_url"]
    assert "app.js" in home["script_src"]


def test_audit_over_screamingfrog_folder(settings, tmp_path):
    _write_internal(tmp_path)
    _write_outlinks(tmp_path)
    summary = run_audit_from_folder(str(tmp_path), settings, settings.audits_dir)

    assert summary["source"] == "screamingfrog"
    assert summary["urls_assessed"] == 7
    assert summary["total_checks"] == 91
    st = _status_map(summary)

    # Checks the SF Internal export supports -> real verdicts over all URLs.
    assert st[34] == "Present"   # missing H1
    assert st[79] == "Present"   # missing canonical
    assert st[76] == "Present"   # multiple titles (Title 2)
    assert st[56] == "Present"   # multiple meta descriptions
    assert st[61] == "Present"   # nosnippet
    assert st[62] == "Present"   # noarchive
    assert st[63] == "Present"   # noindex + canonical
    # Link-graph checks supported via All Outlinks.
    assert st[18] == "Present"   # utm in internal links
    assert st[66] == "Present"   # nofollow internal link to indexable page

    # Checks whose data the export does NOT provide -> Not assessed, not false.
    for num in (7, 23, 33, 52, 30, 31, 32, 41, 80):
        assert st[num] == "Not assessed", f"#{num} should be Not assessed on SF data"

    # No illegal states anywhere.
    assert set(st.values()) <= {"Present", "Not present", "Not assessed", "Heuristic"}


def test_internal_only_gates_link_checks(settings, tmp_path):
    _write_internal(tmp_path)  # no outlinks export
    summary = run_audit_from_folder(str(tmp_path), settings, settings.audits_dir)
    st = _status_map(summary)
    # Column-based checks still work...
    assert st[34] == "Present"
    assert st[76] == "Present"
    # ...but link-graph checks are now Not assessed (no outlinks file).
    for num in (18, 49, 64, 65, 66, 91):
        assert st[num] == "Not assessed", f"#{num} needs links"


def test_gate_note_names_the_missing_signal(settings, tmp_path):
    _write_internal(tmp_path)
    summary = run_audit_from_folder(str(tmp_path), settings, settings.audits_dir)
    from openpyxl import load_workbook

    wb = load_workbook(summary["audit_xlsx"], read_only=True, data_only=True)
    notes = {int(r[0]): r[8] for i, r in enumerate(wb["Checklist"].iter_rows(values_only=True))
             if i > 0 and r[0] is not None}
    wb.close()
    assert "viewport" in (notes[52] or "").lower()
    assert "screamingfrog" in (notes[52] or "").lower()


def test_missing_internal_export_errors(tmp_path):
    # A folder with a CSV that has no Address column is not a usable SF export.
    pd.DataFrame([{"foo": "bar"}]).to_csv(tmp_path / "random.csv", index=False)
    with pytest.raises(ValueError):
        ingest_folder(str(tmp_path))


def test_empty_folder_errors(tmp_path):
    with pytest.raises(ValueError):
        ingest_folder(str(tmp_path))
