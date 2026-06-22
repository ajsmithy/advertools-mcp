"""Load the audit check catalogue (xlsx) — the single source of truth for checks."""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

CATALOGUE_FILENAME = "advertools_audit_check_catalogue.xlsx"
DEFAULT_CATALOGUE_PATH = Path(__file__).with_name(CATALOGUE_FILENAME)

VALID_TIERS = {"CRAWL", "CRAWL+", "RENDER", "CWV", "EXTERNAL"}


@dataclass(frozen=True)
class CheckDef:
    num: int
    check: str
    tier: str
    detect: str
    extra: str
    verdict: str


@lru_cache(maxsize=4)
def load_catalogue(path: str | None = None) -> tuple[CheckDef, ...]:
    src = Path(path) if path else DEFAULT_CATALOGUE_PATH
    if not src.exists():
        raise FileNotFoundError(f"Audit catalogue not found: {src}")
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb["Audit Checks"]
    checks: list[CheckDef] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or row[0] is None:  # header / blank
            continue
        checks.append(
            CheckDef(
                num=int(row[0]),
                check=str(row[1]).strip(),
                tier=str(row[2]).strip(),
                detect=str(row[3]).strip() if row[3] else "",
                extra=str(row[4]).strip() if row[4] else "",
                verdict=str(row[5]).strip() if row[5] else "",
            )
        )
    wb.close()
    return tuple(checks)
