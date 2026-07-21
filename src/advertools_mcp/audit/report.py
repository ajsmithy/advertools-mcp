"""Write the audit workbook: checklist, detail, and summary sheets."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_STATUS_FILL = {
    "Present": "FFF4CCCC",       # finding present (issue)
    "Not present": "FFD9EAD3",   # clean
    "Heuristic": "FFFCE5CD",     # heuristic finding
    "Not assessed": "FFEFEFEF",  # could not evaluate
}
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="FF374151")


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _autosize(ws, max_width: int = 80) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, max_width)


def write_report(
    path: Path,
    results: list[dict[str, Any]],
    detail_rows: list[dict[str, Any]],
    config,
    crawl_detail=None,
) -> None:
    wb = Workbook()

    # ---- Checklist sheet ----
    ws = wb.active
    ws.title = "Checklist"
    headers = ["#", "Check", "Tier", "Status", "What it checks & why it matters",
               "Affected URL count", "Example URLs", "Detection source", "Note"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for r in results:
        ws.append(
            [
                r["num"], r["check"], r["tier"], r["status"], r.get("description", ""),
                r["affected_count"], "\n".join(r["example_urls"]), r["detection_source"], r["note"],
            ]
        )
        fill = _STATUS_FILL.get(r["status"])
        if fill:
            ws.cell(row=ws.max_row, column=4).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=ws.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    _autosize(ws)
    ws.column_dimensions["E"].width = 70  # description column reads better fixed-width

    # ---- Crawl Detail sheet (the flat crawl-detail export, embedded) ----
    if crawl_detail is not None and len(crawl_detail):
        wc = wb.create_sheet("Crawl Detail")
        wc.append(list(crawl_detail.columns))
        _style_header(wc, len(crawl_detail.columns))
        for record in crawl_detail.itertuples(index=False, name=None):
            wc.append([("" if v is None else v) for v in record])
        wc.freeze_panes = "B2"
        _autosize(wc, max_width=60)

    # ---- Detail sheet (one row per offending URL) ----
    wd = wb.create_sheet("Detail")
    dheaders = ["#", "Check", "Tier", "Offending URL"]
    wd.append(dheaders)
    _style_header(wd, len(dheaders))
    for d in detail_rows:
        wd.append([d["num"], d["check"], d["tier"], d["url"]])
    wd.freeze_panes = "A2"
    _autosize(wd)

    # ---- Summary sheet ----
    wsum = wb.create_sheet("Summary")
    wsum.append(["Status", "Count"])
    _style_header(wsum, 2)
    by_status: dict[str, int] = {}
    by_tier: dict[str, dict[str, int]] = {}
    for r in results:
        by_status[r["status"]] = by_status.get(r["status"], 0) + 1
        tier_counts = by_tier.setdefault(r["tier"], {})
        tier_counts[r["status"]] = tier_counts.get(r["status"], 0) + 1
    for status, count in sorted(by_status.items()):
        wsum.append([status, count])

    wsum.append([])
    trow = wsum.max_row + 1
    wsum.cell(row=trow, column=1, value="Tier").font = Font(bold=True)
    wsum.cell(row=trow, column=2, value="Status").font = Font(bold=True)
    wsum.cell(row=trow, column=3, value="Count").font = Font(bold=True)
    for tier, statuses in sorted(by_tier.items()):
        for status, count in sorted(statuses.items()):
            wsum.append([tier, status, count])

    wsum.append([])
    crow = wsum.max_row + 1
    wsum.cell(row=crow, column=1, value="Audit configuration").font = Font(bold=True)
    for key, value in config.as_rows():
        wsum.append([key, value])
    _autosize(wsum)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
